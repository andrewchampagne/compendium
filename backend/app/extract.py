"""Text extraction for uploaded documents (pptx, docx, pdf, txt/md).

Documents are split into segments (slide, page, or section) so cards can link back to where they came from.
"""
import io
import re
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED = {".pptx", ".docx", ".pdf", ".txt", ".md"}

DOCX_PARAS_PER_SECTION = 12
TEXT_CHARS_PER_SECTION = 1500

Segments = list[tuple[str, str]]  # (loc, text), e.g. ("Slide 3", "...")


def _pptx(data: bytes) -> Segments:
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    lines.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[notes] {notes}")
        if lines:
            out.append((f"Slide {i}", "\n".join(lines)))
    return out


def _docx(data: bytes) -> Segments:
    doc = Document(io.BytesIO(data))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    has_headings = any((p.style.name or "").startswith("Heading") for p in paras)
    sections: list[list[str]] = []
    for p in paras:
        is_heading = (p.style.name or "").startswith("Heading")
        start_new = (
            not sections
            or (has_headings and is_heading)
            or (not has_headings and len(sections[-1]) >= DOCX_PARAS_PER_SECTION)
        )
        if start_new:
            sections.append([])
        sections[-1].append(p.text)
    table_lines = [
        " | ".join(c.text.strip() for c in row.cells) for table in doc.tables for row in table.rows
    ]
    if table_lines:
        if not sections:
            sections.append([])
        sections[-1].extend(table_lines)
    return [(f"Section {i}", "\n".join(s)) for i, s in enumerate(sections, 1)]


def _pdf(data: bytes) -> Segments:
    reader = PdfReader(io.BytesIO(data))
    out = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            out.append((f"Page {i}", text))
    return out


def _plain(text: str) -> Segments:
    if re.search(r"^#{1,6}\s", text, re.M):
        parts = re.split(r"(?m)^(?=#{1,6}\s)", text)
    else:
        parts, buf = [], ""
        for block in re.split(r"\n\s*\n", text):
            if buf and len(buf) + len(block) > TEXT_CHARS_PER_SECTION:
                parts.append(buf)
                buf = ""
            buf = f"{buf}\n\n{block}" if buf else block
        if buf:
            parts.append(buf)
    parts = [p.strip() for p in parts if p.strip()]
    return [(f"Section {i}", p) for i, p in enumerate(parts, 1)]


def extract_segments(filename: str, data: bytes) -> Segments:
    ext = Path(filename).suffix.lower()
    if ext == ".pptx":
        return _pptx(data)
    if ext == ".docx":
        return _docx(data)
    if ext == ".pdf":
        return _pdf(data)
    if ext in (".txt", ".md"):
        return _plain(data.decode("utf-8", errors="replace"))
    raise ValueError(f"Unsupported file type: {ext}")


def extract_text(filename: str, data: bytes) -> str:
    return "\n\n".join(f"## {loc}\n{text}" for loc, text in extract_segments(filename, data))
